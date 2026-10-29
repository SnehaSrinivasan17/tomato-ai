"""Section-aware DOCX extraction, provenance checks and idempotent Qdrant upsert."""
import csv
import hashlib
import json
import re
import uuid
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from src.rag.settings import POLICIES, RESULTS, COLLECTION, setting

SECTION = re.compile(r'^([A-Z]{2,}(?:-[A-Z]+)?-\d{2,3})\s*\|\s*(.+)$')


def blocks(document):
    for child in document.element.body:
        if child.tag.endswith('}p'):
            text = Paragraph(child, document).text.strip()
            if text:
                yield text
        elif child.tag.endswith('}tbl'):
            table = Table(child, document)
            for row in table.rows:
                yield ' | '.join(cell.text.strip() for cell in row.cells)


def rows(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as file:
        return list(csv.DictReader(file))


def load_chunks(directory=RESULTS):
    return [json.loads(line) for line in (Path(directory)/'rag_chunks.jsonl').read_text(encoding='utf8').splitlines() if line.strip()]


def prepare(directory=POLICIES, output=RESULTS):
    directory, output = Path(directory), Path(output)
    output.mkdir(parents=True, exist_ok=True)
    manifest_rows = rows(directory/'retrieval_manifest.csv')
    manifest = {row['policy_id']:row for row in manifest_rows}
    chunks, duplicates, mismatches, seen = [], [], [], set()
    documents = []
    def finish(filename, metadata, section, lines):
        if not section:
            return
        policy_id, title = section
        if policy_id in seen:
            duplicates.append(policy_id)
        seen.add(policy_id)
        fields = {}
        for line in lines:
            key, sep, value = line.partition(':')
            if sep:
                fields[key.strip().lower().replace(' ', '_')] = value.strip()
        control = fields.get('control_metadata','')
        severity = re.search(r'Severity\s+(P[0-3](?:/P[0-3])?)',control)
        authority = re.search(r'Authority\s+([^|]+)',control)
        source_ids = re.findall(r'SRC-\d+',control)
        row = dict(policy_id=policy_id,section_id=policy_id,title=title,**metadata,
            rule_text=fields.get('rule',''),mandatory_actions=fields.get('mandatory_action'),
            permitted_actions=fields.get('permitted_action'),prohibited_actions=fields.get('prohibited_action'),
            evidence_requirements=fields.get('evidence'),owner=fields.get('owner'),escalation=fields.get('escalation'),
            exceptions=fields.get('exception'),applicable_conditions=fields.get('trigger'),severity=severity.group(1) if severity else None,
            authority=authority.group(1).strip() if authority else None,source_ids=source_ids,
            source_filename=filename,kind='template' if policy_id.startswith('TMP-') else 'policy',
            simulation=True,approved_by=None,effective_date=None,active=None,review_date=None,
            approval_requirements='Requires Tomato approval; draft only',topic=title,
            chunk_text=f'{policy_id} | {title}\n'+'\n'.join(lines))
        reference = manifest.get(policy_id)
        row['jurisdiction'] = reference.get('jurisdiction') if reference else None
        for key in ['doc_id','title','version','severity','owner','authority']:
            if reference and str(row.get(key)) != str(reference.get(key)):
                mismatches.append(dict(policy_id=policy_id,field=key,document=row.get(key),manifest=reference.get(key)))
        if reference and reference['status'] != row['approval_status']:
            mismatches.append(dict(policy_id=policy_id,field='status',document=row['approval_status'],manifest=reference['status']))
        row['chunk_id'] = str(uuid.uuid5(uuid.NAMESPACE_URL,f'tomato/{filename}/{policy_id}/{row["version"]}'))
        row['content_sha256'] = hashlib.sha256(row['chunk_text'].encode()).hexdigest()
        chunks.append(row)
    for path in sorted(directory.glob('*.docx')):
        content = list(blocks(Document(path)))
        header = next((line for line in content if line.startswith('POLICY ID:')), '')
        doc_id = re.search(r'POLICY ID:\s*(\S+)',header)
        version = re.search(r'VERSION:\s*([^|]+)',header)
        if not doc_id or not version:
            raise ValueError('Document metadata missing: '+path.name)
        status = 'PROPOSED_REQUIRES_APPROVAL' if 'PROPOSED' in header else None
        preamble = content[:next((i for i,text in enumerate(content) if SECTION.match(text)),len(content))]
        metadata = dict(doc_id=doc_id.group(1),document_title=content[0],version=version.group(1).strip(),
            approval_status=status,source_status=header,document_controls='\n'.join(preamble))
        section, lines = None, []
        for text in content:
            match = SECTION.match(text)
            if match:
                finish(path.name,metadata,section,lines)
                section, lines = match.groups(), []
            elif text == 'Source and approval notice':
                finish(path.name,metadata,section,lines)
                section,lines=None,[]
            elif section:
                lines.append(text)
        finish(path.name,metadata,section,lines)
        documents.append(dict(filename=path.name,sha256=hashlib.sha256(path.read_bytes()).hexdigest(),tables=len(Document(path).tables)))
    section_ids = {chunk['policy_id'] for chunk in chunks}
    scenarios = rows(directory/'evaluation_scenarios.csv')
    source_ids = {row['source_id'] for row in rows(directory/'source_register.csv')}
    conflicts = rows(directory/'policy_conflict_matrix.csv')
    # An actual contradiction, preserved verbatim; precedence is not guessed.
    if {'SM-102','SM-SLA-01'}.issubset(section_ids):
        conflicts.append(dict(rule_a='SM-102',rule_b='SM-SLA-01',precedence='',
            reason='Unresolved proposed SLA: P2 60 minutes vs 2 hours; P3 4 vs 8 staffed hours.',unresolved=True))
    templates = [json.loads(line) for line in (directory/'approved_templates.jsonl').read_text(encoding='utf8').splitlines() if line.strip()]
    for template in templates:
        # Filename is not proof of approval. No approver, version or active date is invented.
        text = template['scenario']+'\n'+template['public_response']+'\n'+template['private_followup']
        chunks.append(dict(chunk_id=str(uuid.uuid5(uuid.NAMESPACE_URL,'tomato/template/'+template['template_id'])),
            policy_id=template['policy_id'],template_id=template['template_id'],doc_id=None,title=template['scenario'],
            rule_text=template['public_response'],chunk_text=text,source_filename='approved_templates.jsonl',version=None,
            approval_status=None,source_status='Approval not supplied in template record',approved_by=None,active=None,
            effective_date=None,severity=template.get('severity'),authority=None,jurisdiction=None,topic=template['scenario'],
            kind='template',simulation=True,template=template,content_sha256=hashlib.sha256(text.encode()).hexdigest()))
    report = dict(documents=documents,document_sections=len(section_ids),total_records=len(chunks),
        policy_records=sum(c['kind']=='policy' for c in chunks),template_records=sum(c['kind']=='template' for c in chunks),
        duplicate_ids=duplicates,duplicate_manifest_ids=len(manifest_rows)-len(manifest),
        missing_manifest_ids=sorted(section_ids-set(manifest)),missing_document_ids=sorted(set(manifest)-section_ids),
        metadata_mismatches=mismatches,missing_expected_ids=sorted({r['expected_policy_id'] for r in scenarios}-section_ids),
        missing_source_ids=sorted({s for c in chunks for s in c.get('source_ids',[])}-source_ids),
        conflicts=conflicts,unresolved_conflict_ids=sorted({r[k] for r in conflicts for k in ['rule_a','rule_b'] if r.get(k) not in section_ids}),
        approval='All document policies proposed; template JSONL lacks approval metadata.',
        hierarchy=(directory/'policy_hierarchy.md').read_text(encoding='utf8'),scenario_count=len(scenarios),
        approved_operational_records=0)
    if duplicates or report['missing_document_ids'] or mismatches:
        report['ingestion_blocked'] = True
    else:
        report['ingestion_blocked'] = False
    (output/'rag_chunks.jsonl').write_text('\n'.join(json.dumps(c,ensure_ascii=False) for c in chunks)+'\n',encoding='utf8')
    (output/'policy_audit.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf8')
    crosswalk = output/'policy_id_crosswalk.csv'
    if not crosswalk.exists():
        with crosswalk.open('w',encoding='utf8',newline='') as file:
            writer=csv.DictWriter(file,fieldnames=['expected_policy_id','mapped_policy_id','reviewed_by','reviewed_at','rationale'])
            writer.writeheader()
            for policy_id in report['missing_expected_ids']:
                writer.writerow(dict(expected_policy_id=policy_id,rationale='Missing: requires human review; no automatic mapping.'))
    return chunks, report


def qdrant_client():
    from qdrant_client import QdrantClient
    url,key=setting('QDRANT_URL'),setting('QDRANT_API_KEY')
    if not url or not key:
        raise RuntimeError('QDRANT_URL and QDRANT_API_KEY are required for Qdrant Cloud.')
    return _cloud_client(url,key)


from functools import lru_cache
@lru_cache(maxsize=2)
def _cloud_client(url,key):
    from qdrant_client import QdrantClient
    return QdrantClient(url=url,api_key=key,timeout=30)


def ingest(client=None, encoder=None):
    from qdrant_client import models
    from src.rag.embeddings import BGEEncoder
    chunks,audit=prepare()
    if audit['ingestion_blocked']:
        raise RuntimeError('Policy audit blocks ingestion; review policy_audit.json.')
    client=client or qdrant_client()
    encoder=encoder or BGEEncoder()
    chunks=[c for c in chunks if c['kind']=='policy']
    vectors=encoder.encode([c['chunk_text'] for c in chunks],batch_size=1)
    import numpy as np
    if vectors.shape != (len(chunks),1024) or not np.isfinite(vectors).all():
        raise RuntimeError('BGE-M3 requires finite 1024-dimensional vectors.')
    dimension=int(vectors.shape[1])
    if not client.collection_exists(COLLECTION):
        client.create_collection(COLLECTION,vectors_config=models.VectorParams(size=dimension,distance=models.Distance.COSINE))
    info=client.get_collection(COLLECTION)
    if info.config.params.vectors.size != dimension or str(info.config.params.vectors.distance).lower() != 'cosine':
        raise RuntimeError('Collection vector size/distance does not match actual BGE-M3 output.')
    # Qdrant Cloud strict mode requires indexes for every filtered payload field.
    for field in ['ingestion_owner','kind','embedding_model','policy_id','source_filename']:
        if field not in info.payload_schema:
            client.create_payload_index(COLLECTION,field_name=field,field_schema=models.PayloadSchemaType.KEYWORD,wait=True)
    for start in range(0,len(chunks),16):
        client.upsert(COLLECTION,points=[models.PointStruct(id=c['chunk_id'],vector=v.tolist(),payload=dict(c,
            embedding_model=encoder.model_name,ingestion_owner='tomato-policy-rag')) for c,v in zip(chunks[start:start+16],vectors[start:start+16])],wait=True)
    # Synchronize only stale points explicitly owned by this policy ingester.
    old,_=client.scroll(COLLECTION,scroll_filter=models.Filter(must=[models.FieldCondition(key='ingestion_owner',match=models.MatchValue(value='tomato-policy-rag'))]),limit=10000)
    current={c['chunk_id'] for c in chunks}
    stale=[p.id for p in old if str(p.id) not in current]
    if stale:
        client.delete(COLLECTION,points_selector=models.PointIdsList(points=stale),wait=True)
    revision_path=RESULTS/'models/models--BAAI--bge-m3/refs/main'
    source_revision=revision_path.read_text().strip() if revision_path.exists() else None
    result=dict(embedding_source_revision=source_revision,collection=COLLECTION,documents=len({c['source_filename'] for c in chunks}),records=len(chunks),dimension=dimension,embedding_model=encoder.model_name,
                stale_owned_points_removed=len(stale),total_collection_points=client.count(COLLECTION,exact=True).count)
    RESULTS.mkdir(exist_ok=True)
    (RESULTS/'ingestion.json').write_text(json.dumps(result,indent=2))
    return result


def parse_policy_file(path):
    """Supported input adapters; only explicit policy IDs become sections."""
    path=Path(path)
    if path.suffix.lower()=='.docx': lines=list(blocks(Document(path)))
    elif path.suffix.lower() in {'.md','.txt'}: lines=path.read_text(encoding='utf-8-sig').splitlines()
    elif path.suffix.lower() in {'.jsonl','.csv'}:
        records=rows(path) if path.suffix.lower()=='.csv' else [json.loads(x) for x in path.read_text(encoding='utf-8-sig').splitlines() if x.strip()]
        return [dict(r,source_filename=path.name) for r in records if r.get('policy_id') and (r.get('rule_text') or r.get('policy_text'))]
    else: raise ValueError('Unsupported policy source format.')
    sections=[];current=None
    for line in lines:
        match=SECTION.match(line.lstrip('# ').strip())
        if match:
            current=dict(policy_id=match.group(1),title=match.group(2),source_filename=path.name,policy_text='')
            sections.append(current)
        elif current:current['policy_text']+=line+'\n'
    return sections
